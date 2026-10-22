from pathlib import Path
import fitz
from pptx import Presentation

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / 'data' / 'homework_previews'
JOBS = [
    ('mc_ch8_1', 'course_materials/现代控制理论/homework/1199059/第8章-作业1.pdf'),
    ('mc_ch8_2', 'course_materials/现代控制理论/homework/1199436/HW_5.pdf'),
    ('pc_week7', 'course_materials/过程控制工程/homework/1200101/作业题(1).pptx'),
]

def extract_shape(shape):
    vals = []
    if shape.has_text_frame:
        vals.extend(p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip())
    if shape.has_table:
        for row in shape.table.rows:
            vals.append(' | '.join(c.text.strip().replace('\n', ' ') for c in row.cells))
    if shape.shape_type == 6:
        for sub in shape.shapes:
            vals.extend(extract_shape(sub))
    return vals

for label, location in JOBS:
    source = ROOT / location
    work = OUT / label
    work.mkdir(parents=True, exist_ok=True)
    md = ['# Extracted source text', f'Source: {location}',
          'Extraction generated mechanically. Check formulas and images against page previews.']
    if source.suffix.lower() == '.pdf':
        doc = fitz.open(source)
        md.append(f'Pages: {len(doc)}')
        for i, page in enumerate(doc):
            md.append(f'\n## Page {i + 1}')
            md.append(page.get_text('text', sort=True))
            page.get_pixmap(matrix=fitz.Matrix(1.3, 1.3), alpha=False).save(str(work / f'page_{i + 1:02}.png'))
    else:
        deck = Presentation(str(source))
        md.append(f'Slides: {len(deck.slides)}')
        for i, slide in enumerate(deck.slides):
            md.append(f'\n## Slide {i + 1}')
            for j, shape in enumerate(slide.shapes):
                vals = extract_shape(shape)
                if vals:
                    md.append(f'### Shape {j}\n' + '\n'.join(vals))
                if shape.shape_type == 13:
                    image = shape.image
                    name = f'slide_{i+1:02}_image_{j:02}.{image.ext}'
                    (work / name).write_bytes(image.blob)
                    md.append(f'Image: {name} ({image.size} px)')
    (work / 'questions.md').write_text('\n\n'.join(md), encoding='utf-8')
    print(f'Extracted {label}: {len(md)} text entries')
